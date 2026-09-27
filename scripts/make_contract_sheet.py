"""Build the Task 19 manual-collection sheet: contract start/end for the top-N players.

One row per player: the N most valuable players today (latest Transfermarkt value) that
have a usable panel row. Each row is pinned to one past season; collectors look up the
contract that was active at that season's 1 July anchor.

Kept as the record of how the 500 players were chosen (2026-09-23). It reads the
retired Kaggle panel (data/processed/panel.parquet), so rerunning it needs PR #2's pipeline
from git history; see docs/scope_change.md.

Usage:
    python scripts/make_contract_sheet.py --data-dir <Dataset root> --out contract_collection.xlsx
"""
import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

COLLECTORS = ("Yunus", "Jack")
# Seasons after this mostly have the same contract as today's snapshot, so they add little.
LATEST_SEASON = 2023

PREFILLED = ["rank", "collector", "player_id", "player_name", "season", "anchor_date",
             "club_that_season", "value_today_eur", "hint_current_contract_end", "hint_joined_current_club"]
TO_FILL = ["contract_start", "contract_end", "source_url", "confidence", "notes"]
COLUMNS = PREFILLED + TO_FILL


def build_rows(data_dir: Path, panel_path: Path, n: int, seed: int) -> pd.DataFrame:
    players = pd.read_parquet(data_dir / "player-scores/players.parquet")
    clubs = pd.read_parquet(data_dir / "player-scores/clubs.parquet", columns=["club_id", "name"])
    profiles = pd.read_parquet(data_dir / "football-datasets/player_profiles.parquet",
                               columns=["player_id", "joined"])
    panel = pd.read_parquet(panel_path, columns=["player_id", "club_id", "season", "y_h1", "is_primary"])

    eligible = panel[(panel.season <= LATEST_SEASON) & panel.y_h1.notna() & panel.is_primary]
    top = (players[players.player_id.isin(eligible.player_id)]
           .sort_values("market_value_in_eur", ascending=False).head(n).reset_index(drop=True))

    # Seeded random season per player so the rows spread across eras instead of piling up in one year.
    rng = np.random.default_rng(seed)
    picked = (eligible[eligible.player_id.isin(top.player_id)]
              .groupby("player_id", group_keys=False)
              .apply(lambda g: g.iloc[[rng.integers(len(g))]]))

    out = (top[["player_id", "name", "market_value_in_eur", "contract_expiration_date"]]
           .merge(picked[["player_id", "season", "club_id"]], on="player_id")
           .merge(clubs.rename(columns={"name": "club_that_season"}), on="club_id", how="left")
           .merge(profiles.drop_duplicates("player_id"), on="player_id", how="left"))
    out["rank"] = np.arange(1, len(out) + 1)
    # Alternate by rank so neither collector gets only the most (or least) famous players.
    out["collector"] = np.where(out["rank"] % 2 == 1, COLLECTORS[0], COLLECTORS[1])
    out["anchor_date"] = pd.to_datetime(out.season.astype(str) + "-07-01").dt.date
    out["season"] = out.season.astype(str) + "/" + (out.season + 1).astype(str).str[-2:]
    out["hint_current_contract_end"] = pd.to_datetime(out.contract_expiration_date, errors="coerce").dt.date
    out["hint_joined_current_club"] = pd.to_datetime(out.joined, errors="coerce").dt.date
    out = out.rename(columns={"name": "player_name", "market_value_in_eur": "value_today_eur"})
    for col in TO_FILL:
        out[col] = None
    return out[COLUMNS]


INSTRUCTIONS = [
    "Contract data collection - Task 19 (manual dataset)",
    "",
    "Goal: for each row, find the contract the player had with 'club_that_season' on 'anchor_date' (1 July).",
    "Filter the 'collector' column to see your own rows (Yunus = odd ranks, Jack = even ranks).",
    "",
    "Columns to fill (white):",
    "  contract_start - date that contract started (or was signed). YYYY-MM-DD. If only the year is known, use YYYY-07-01.",
    "  contract_end   - date that contract was due to end. Most end 30 June, so a year-only answer becomes YYYY-06-30.",
    "  source_url     - REQUIRED. Club announcement, news article or Wikipedia page that states the end date.",
    "  confidence     - high = official/club source; med = reputable news or Wikipedia; low = inferred.",
    "  notes          - anything odd (loan, option year, contract extended during that season, etc.).",
    "",
    "Rules:",
    "  * If the contract was EXTENDED before the anchor date, enter the extension (signed date -> new end).",
    "  * If the player was ON LOAN that season, enter the contract with the parent club and write 'loan' in notes.",
    "  * Grey columns are prefilled - do not edit them. hint_* columns are today's contract/join date; they help",
    "    only when the player has not moved or re-signed since that season.",
    "  * If you cannot find it in ~5 minutes, leave the dates blank, write 'not found' in notes, and move on.",
]


def write_xlsx(rows: pd.DataFrame, out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "contracts"
    ws.append(COLUMNS)
    for rec in rows.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in rec])
    last = len(rows) + 1

    def letter(name):
        return ws.cell(1, COLUMNS.index(name) + 1).column_letter

    def cells(name):
        idx = COLUMNS.index(name) + 1
        return (c for (c,) in ws.iter_rows(min_row=2, max_row=last, min_col=idx, max_col=idx))

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    grey = PatternFill("solid", fgColor="E7E6E6")
    for name in PREFILLED:
        for cell in cells(name):
            cell.fill = grey
    for cell in cells("value_today_eur"):
        cell.number_format = "#,##0"
    for name in ("anchor_date", "hint_current_contract_end", "hint_joined_current_club",
                 "contract_start", "contract_end"):
        for cell in cells(name):
            cell.number_format = "yyyy-mm-dd"

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{letter(COLUMNS[-1])}{last}"
    widths = {"player_name": 24, "club_that_season": 26, "source_url": 40, "notes": 30}
    for name in COLUMNS:
        ws.column_dimensions[letter(name)].width = widths.get(name, 14)

    dates = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                           formula2="DATE(2040,12,31)", allow_blank=True, showErrorMessage=True,
                           error="Enter a date (YYYY-MM-DD) between 2000 and 2040.")
    conf = DataValidation(type="list", formula1='"high,med,low"', allow_blank=True)
    who = DataValidation(type="list", formula1=f'"{",".join(COLLECTORS)}"')
    for dv in (dates, conf, who):
        ws.add_data_validation(dv)
    dates.add(f"{letter('contract_start')}2:{letter('contract_end')}{last}")
    conf.add(f"{letter('confidence')}2:{letter('confidence')}{last}")
    who.add(f"{letter('collector')}2:{letter('collector')}{last}")

    info = wb.create_sheet("instructions")
    for line in INSTRUCTIONS:
        info.append([line])
    info.column_dimensions["A"].width = 120
    info["A1"].font = Font(bold=True, size=14)
    wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--data-dir", type=Path, required=True)
    ap.add_argument("--panel", type=Path, default=Path("data/processed/panel.parquet"))
    ap.add_argument("--out", type=Path, default=Path("data/manual/contract_collection.xlsx"))
    ap.add_argument("-n", type=int, default=500)
    ap.add_argument("--seed", type=int, default=0)
    args = ap.parse_args()

    rows = build_rows(args.data_dir, args.panel, args.n, args.seed)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx(rows, args.out)
    print(f"wrote {len(rows)} rows -> {args.out}")
    print(rows.collector.value_counts().to_string())
    print("seasons:", rows.season.value_counts().sort_index().to_dict())


if __name__ == "__main__":
    main()
